import io
import pytest
from docx import Document
from app.services.extractor import ExtractionError, extract_text, extract_text_from_docx, extract_text_from_txt


def test_extract_text_from_txt():
    sample_text = "John Doe\nSoftware Engineer\nPython, FastAPI, Docker"
    content = sample_text.encode("utf-8")
    result = extract_text_from_txt(content)
    assert "John Doe" in result
    assert "FastAPI" in result


def test_extract_text_from_docx():
    doc = Document()
    doc.add_paragraph("Alice Smith")
    doc.add_paragraph("Senior Data Scientist")
    doc.add_paragraph("Skills: Machine Learning, PyTorch, SQL")
    
    file_stream = io.BytesIO()
    doc.save(file_stream)
    content = file_stream.getvalue()

    result = extract_text_from_docx(content)
    assert "Alice Smith" in result
    assert "Senior Data Scientist" in result
    assert "Machine Learning" in result


def test_extract_text_dispatch_unsupported():
    with pytest.raises(ExtractionError, match="Unsupported file format"):
        extract_text(b"malicious content", "malware.exe")


def test_extract_text_empty_content():
    with pytest.raises(ExtractionError, match="No extractable text found"):
        extract_text(b"   \n\n  ", "empty.txt")
